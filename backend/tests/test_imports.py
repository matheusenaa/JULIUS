"""Importação de extratos, exportações e exclusão de conta."""

import io

from openpyxl import load_workbook

OFX = b"""OFXHEADER:100
<OFX><BANKMSGSRSV1><STMTTRNRS><STMTRS><BANKTRANLIST>
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20261001120000<TRNAMT>-45.90<FITID>A1<MEMO>SUPERMERCADO EXTRA
<STMTTRN><TRNTYPE>CREDIT<DTPOSTED>20261002<TRNAMT>3200.00<FITID>A2<MEMO>SALARIO EMPRESA
<STMTTRN><TRNTYPE>DEBIT<DTPOSTED>20261003<TRNAMT>-120,00<FITID>A3<MEMO>POSTO SHELL
</BANKTRANLIST></STMTRS></STMTTRNRS></BANKMSGSRSV1></OFX>"""

CSV = "Data;Descrição;Valor\n01/10/2026;Padaria;-12,50\n02/10/2026;Pix recebido;1.000,00\n".encode("cp1252")


def account(api):
    return api.post("/api/accounts", {"name": "Itaú", "kind": "checking"}).json()


def preview(api, content, name, acc_id):
    r = api.c.post(
        "/api/import/statement/preview",
        files={"file": (name, content)},
        data={"account_id": acc_id},
        headers=api._h(),
    )
    assert r.status_code == 200, r.text
    return r.json()


def test_ofx_import_preview_confirm_and_no_duplicates(api, ids):
    acc = account(api)
    p = preview(api, OFX, "extrato.ofx", acc["id"])
    assert p["format"] == "ofx" and len(p["items"]) == 3 and p["duplicates"] == 0
    by_desc = {i["description"]: i for i in p["items"]}
    assert (
        by_desc["SUPERMERCADO EXTRA"]["amount_cents"] == 4590
        and by_desc["SUPERMERCADO EXTRA"]["type"] == "expense"
    )
    assert by_desc["SALARIO EMPRESA"]["type"] == "income"
    assert by_desc["POSTO SHELL"]["category_id"] == ids["cats"]["Combustível"]
    assert (
        api.get("/api/transactions", params={"account_id": acc["id"]}).json()["total"] == 0
    )  # prévia não grava

    r = api.post("/api/import/statement/confirm", {"account_id": acc["id"], "items": p["items"]})
    assert r.json() == {"created": 3, "skipped": 0}
    again = preview(api, OFX, "extrato.ofx", acc["id"])
    assert again["duplicates"] == 3
    r = api.post("/api/import/statement/confirm", {"account_id": acc["id"], "items": again["items"]})
    assert r.json()["created"] == 0
    balance = next(a for a in api.get("/api/accounts").json() if a["id"] == acc["id"])["balance_cents"]
    assert balance == 320000 - 4590 - 12000


def test_csv_import_brazilian_format(api):
    acc = account(api)
    p = preview(api, CSV, "extrato.csv", acc["id"])
    assert [(i["description"], i["amount_cents"], i["type"]) for i in p["items"]] == [
        ("Padaria", 1250, "expense"),
        ("Pix recebido", 100000, "income"),
    ]


def test_import_rejects_unknown_format(api):
    acc = account(api)
    r = api.c.post(
        "/api/import/statement/preview",
        files={"file": ("x.exe", b"MZ")},
        data={"account_id": acc["id"]},
        headers=api._h(),
    )
    assert r.status_code == 400


def test_export_excel_and_pdf(api):
    acc = account(api)
    api.post(
        "/api/transactions",
        {
            "type": "expense",
            "account_id": acc["id"],
            "amount_cents": 4590,
            "occurred_on": "2026-10-01",
            "description": "Mercado ção",
        },
    )
    x = api.get("/api/export/transactions.xlsx")
    assert x.status_code == 200
    ws = load_workbook(io.BytesIO(x.content))["Lançamentos"]
    assert ws["C2"].value == "Mercado ção" and ws["H2"].value == -45.9
    pdf = api.get("/api/export/report.pdf", params={"month": "2026-10"})
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_delete_account_requires_password_and_removes_everything(api, anon):
    acc = account(api)
    api.post(
        "/api/transactions",
        {
            "type": "expense",
            "account_id": acc["id"],
            "amount_cents": 100,
            "occurred_on": "2026-10-01",
            "description": "x",
        },
    )
    bad = api.post("/api/auth/delete-account", {"password": "errada", "confirm": "EXCLUIR"})
    assert bad.status_code == 400
    ok = api.post("/api/auth/delete-account", {"password": "senha-forte-123", "confirm": "EXCLUIR"})
    assert ok.status_code == 204
    assert api.get("/api/auth/session").json() is None
    r = anon.post("/api/auth/login", {"email": api.email, "password": "senha-forte-123"})
    assert r.status_code == 401
