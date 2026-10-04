"""Exportação em Excel (.xlsx) e relatório mensal em PDF."""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Account, Category, Transaction, User
from app.services import analytics
from app.services.backup import PAYMENT_PT, STATUS_PT, TYPE_PT
from app.services.dates import MONTHS_PT
from app.services.money import brl

BRAND = "0B3D2C"


def transactions_xlsx(db: Session, user_id: str) -> bytes:
    cats = {c.id: c for c in db.scalars(select(Category).where(Category.user_id == user_id))}
    accs = {a.id: a.name for a in db.scalars(select(Account).where(Account.user_id == user_id))}
    wb = Workbook()
    ws = wb.active
    ws.title = "Lançamentos"
    headers = [
        "Data",
        "Tipo",
        "Descrição",
        "Categoria",
        "Subcategoria",
        "Conta",
        "Conta destino",
        "Valor",
        "Situação",
        "Forma de pagamento",
        "Fixa",
        "Observações",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=BRAND)
    for t in db.scalars(
        select(Transaction)
        .where(Transaction.user_id == user_id, Transaction.deleted_at.is_(None))
        .order_by(Transaction.occurred_on, Transaction.created_at)
    ):
        cat = cats.get(t.category_id)
        parent = cats.get(cat.parent_id) if cat and cat.parent_id else None
        signed = t.amount_cents / 100 * (-1 if t.type == "expense" else 1)
        ws.append(
            [
                t.occurred_on,
                TYPE_PT[t.type],
                t.description,
                parent.name if parent else (cat.name if cat else ""),
                cat.name if parent else "",
                accs.get(t.account_id, ""),
                accs.get(t.to_account_id, ""),
                signed,
                STATUS_PT.get(t.status, t.status),
                PAYMENT_PT.get(t.payment_method, ""),
                "Sim" if t.is_fixed else "Não",
                t.notes or "",
            ]
        )
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
        row[7].number_format = "#,##0.00;[Red]-#,##0.00"
    widths = [12, 14, 40, 18, 18, 18, 18, 14, 12, 18, 6, 40]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("Resumo mensal")
    summary.append(["Mês", "Receitas", "Despesas", "Resultado"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    today = date.today()
    for m in analytics.monthly_series(db, user_id, today.replace(day=1), 12):
        summary.append([m["month"], m["income"] / 100, m["expense"] / 100, m["net_cents"] / 100])
    for row in summary.iter_rows(min_row=2):
        row[0].number_format = "MM/YYYY"
        for c in row[1:]:
            c.number_format = "#,##0.00"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _latin(text: str) -> str:
    """Fontes padrão do PDF só têm Latin-1 (acentos do português incluídos)."""
    return text.replace("−", "-").replace("…", "...").encode("latin-1", "replace").decode("latin-1")


def month_report_pdf(db: Session, user: User, ref: date) -> bytes:
    from fpdf import FPDF

    r = analytics.report(db, user.id, "month", ref)
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_page()
    pdf.set_fill_color(11, 61, 44)
    pdf.rect(0, 0, 210, 28, "F")
    pdf.set_text_color(255, 255, 255)
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_xy(12, 9)
    pdf.cell(0, 10, "JULIUS")
    pdf.set_font("Helvetica", "", 11)
    pdf.set_xy(12, 17)
    pdf.cell(0, 8, _latin(f"Relatório de {MONTHS_PT[ref.month - 1]} de {ref.year} - {user.name}"))
    pdf.set_text_color(17, 21, 19)
    pdf.set_xy(12, 36)

    def line(label: str, value: str, bold: bool = False) -> None:
        pdf.set_font("Helvetica", "B" if bold else "", 11)
        pdf.cell(120, 8, _latin(label))
        pdf.cell(0, 8, _latin(value), align="R", new_x="LMARGIN", new_y="NEXT")

    line("Entrou", brl(r["income_cents"]))
    line("Saiu", brl(r["expense_cents"]))
    line("Resultado do mês", brl(r["net_cents"]), bold=True)
    line("Despesas fixas", brl(r["fixed_cents"]))
    line("Despesas variáveis", brl(r["variable_cents"]))
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 9, _latin("Gastos por categoria"), new_x="LMARGIN", new_y="NEXT")
    for c in r["by_category"][:15]:
        line(f"{c['name']} ({round(c['share'] * 100)}%)", brl(c["total_cents"]))
    pdf.ln(4)
    pdf.set_font("Helvetica", "B", 13)
    pdf.cell(0, 9, _latin("Maiores despesas"), new_x="LMARGIN", new_y="NEXT")
    for t in r["top_expenses"]:
        line(f"{t['date'].strftime('%d/%m')}  {t['description'][:60]}", brl(t["amount_cents"]))
    pdf.ln(6)
    pdf.set_font("Helvetica", "I", 9)
    pdf.set_text_color(111, 120, 115)
    pdf.multi_cell(
        0,
        5,
        _latin(
            "Valores calculados pelo JULIUS a partir dos lançamentos realizados. "
            f"Gerado em {date.today().strftime('%d/%m/%Y')}."
        ),
    )
    return bytes(pdf.output())
